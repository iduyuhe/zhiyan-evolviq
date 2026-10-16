# -*- coding: utf-8 -*-
"""生成《造船行业智能体合作方案书》docx —— 零真名铁律合规版本。"""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(BASE, "造船行业智能体合作方案书.docx")

DEEP = RGBColor(0x0B, 0x25, 0x45)   # 主色 深蓝
BRAND = RGBColor(0x25, 0x63, 0xEB)  # 辅色 品牌蓝
CYAN = RGBColor(0x06, 0xB6, 0xD4)   # 强调 科技青
INK = RGBColor(0x0F, 0x17, 0x2A)    # 正文主色
SUB = RGBColor(0x47, 0x55, 0x69)    # 正文次色
FONT = "Microsoft YaHei"

doc = Document()

# ---- 全局默认字体（含东亚字体） ----
def set_run_font(run, size=None, bold=None, color=None):
    run.font.name = FONT
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), FONT)
    rFonts.set(qn('w:ascii'), FONT)
    rFonts.set(qn('w:hAnsi'), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color

# 默认 Normal 样式
normal = doc.styles['Normal']
normal.font.name = FONT
normal.font.size = Pt(11)
normal.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)

# 设置页边距
for s in doc.sections:
    s.top_margin = Cm(2.5)
    s.bottom_margin = Cm(2.5)
    s.left_margin = Cm(2.8)
    s.right_margin = Cm(2.8)

def shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def set_cell_text(cell, text, size=11, bold=False, color=INK, align='left'):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = {'left': WD_ALIGN_PARAGRAPH.LEFT,
                   'center': WD_ALIGN_PARAGRAPH.CENTER,
                   'right': WD_ALIGN_PARAGRAPH.RIGHT}[align]
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)

def add_title(text, size=15, color=DEEP, space_before=6, space_after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    # 左侧色条
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single'); left.set(qn('w:sz'), '24')
    left.set(qn('w:space'), '6'); left.set(qn('w:color'), '2563EB')
    pbdr.append(left)
    pPr.append(pbdr)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=True, color=color)
    return p

def add_para(text, size=11, color=INK, bold=False, space_after=6, indent=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    if indent:
        p.paragraph_format.left_indent = Cm(indent)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p

def add_bullet(text, size=11, color=INK, bold=False):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p

def make_table(headers, rows, col_widths=None, header_fill='0B2545'):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = 'Table Grid'
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_text(hdr[i], h, size=11, bold=True, color=RGBColor(0xFF,0xFF,0xFF), align='center')
        shade_cell(hdr[i], header_fill)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            set_cell_text(cells[i], val, size=10.5, color=INK, align='left')
            if i == 0:
                set_run_font(cells[i].paragraphs[0].runs[0], size=10.5, bold=True, color=DEEP)
    if col_widths:
        for i, w in enumerate(col_widths):
            for r in t.rows:
                r.cells[i].width = Cm(w)
    return t

# ============ 封面 ============
doc.add_paragraph().add_run("").font.size = Pt(40)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("智衍 EvolvIQ"); set_run_font(r, size=40, bold=True, color=DEEP)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("造船行业智能体合作方案书"); set_run_font(r, size=26, bold=True, color=BRAND)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("工业 AI 决策大脑 · 让每艘船拥有可对话的决策分身"); set_run_font(r, size=14, color=SUB)
doc.add_paragraph().add_run("").font.size = Pt(30)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("杜玉河 · 工业 5 点 0 产业生态联盟"); set_run_font(r, size=13, color=SUB)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("2026 年 10 月"); set_run_font(r, size=12, color=SUB)

doc.add_page_break()

# ============ 目录式导言 ============
add_title("一、背景与目的", size=15)
add_para("中国造船业连续多年位居全球前列，但在「多品种、小批量、长周期」的项目制模式下，经营管控者长期面对进度、成本、质量、设备、研发多源信息割裂的困境：决策依赖经验、响应滞后、归因困难。")
add_para("智衍 EvolvIQ 是一套面向工业的 AI 决策平台，以「工业 AI 决策分身」为核心，在客户既有的 ERP / MES 之上叠加实时决策脑与全息真相源，而非推倒重来。本方案书系统梳理造船行业的痛点、智衍的能力矩阵与合作进阶路径，作为与造船企业系统化交流、迈向智能体深度合作的共同基线。")
add_para("说明：出于零真名铁律，本方案书中涉及的客户均作匿名化处理，真实主体仅限内部视图，不对第三方披露。", size=10.5, color=SUB, bold=True)

# ============ 二、行业痛点 ============
add_title("二、造船行业的决策痛点", size=15)
add_para("基于与多家头部造船集团（含演示租户与培训接触客户）的交流，造船企业的决策瓶颈高度集中在以下四类：")
make_table(
    ["#", "决策痛点", "典型表现", "影响"],
    [
        ["1", "多品种·小批量·长周期", "一艘船即一个定制项目，分段、舾装、调试跨数月", "进度、资源、成本难以实时掌控"],
        ["2", "大型专用设备多", "龙门吊、焊接机器人、船坞、船用柴油机停机损失高", "运维依赖老师傅经验，故障被动救火"],
        ["3", "一次合格率低", "船体报验环节多，缺陷归因链条长", "返工代价巨大，质量成本不可控"],
        ["4", "跨专业协同复杂", "设计、工艺、生产、质量数据割裂", "变更影响难量化，协同成本高"],
    ],
    col_widths=[1.0, 4.2, 6.0, 5.0],
)

# ============ 三、智衍定位 ============
add_title("三、智衍定位：工业 AI 决策分身", size=15)
add_para("智衍不替代企业的 ERP / MES，而是在其之上叠加「实时决策脑 + 全息真相源」。其能力以四层决策分身框架组织，与客户的决策层级天然对齐：")
make_table(
    ["层级", "决策分身", "核心职责", "价值落点"],
    [
        ["产业级", "产业决策分身", "全球价值链与宏观周期洞察", "先于行业看清趋势"],
        ["行业级", "行业洞察分身", "对标与行业基准", "找准自身坐标"],
        ["企业级", "经营管控分身", "驾驶舱与经营决策", "实时掌控经营全局"],
        ["专业级", "岗位执行分身", "设备 / 工艺 / 质量执行", "直达现场与设备"],
    ],
    col_widths=[2.2, 3.4, 6.0, 4.6],
)
add_para("成熟度阶梯：L0 观察 → L1 判断 → L2 预案 → L3 授权执行。当前大多数场景处于 L0–L2，对外不承诺 L3——分身提供判断与预案，最终决策与签署由人终审（人留终审红线）。", space_after=4)
add_para("护城河：四层框架由同一套知识资产与 Agent 贯通，支持跨层下钻，而非孤立的工具集合。", size=10.5, color=SUB)

# ============ 四、造船九类智能体矩阵 ============
add_title("四、造船九类智能体矩阵", size=15)
add_para("行业成为数据驱动的一等公民：智能体的名称、场景与设备类型随行业 manifest 自动切换。面向造船业务域，智衍落地九类智能体：")
make_table(
    ["业务域", "智能体", "核心能力"],
    [
        ["设备", "设备健康 Agent", "健康诊断、预测维护建议、关键部件寿命管理"],
        ["装配", "分段装配 Agent", "分段合拢进度、节拍与资源协同"],
        ["质量", "涂装焊接缺陷判定 Agent", "缺陷误报过滤、阈值优化、复判工时节省"],
        ["质量", "一次合格率 Agent", "报验合格率趋势、缺陷分类、设备-质量根因"],
        ["质量", "质量追溯 Agent", "跨工序质量问题归因与闭环"],
        ["工艺", "可建造性检查 Agent", "焊接可达性 / 精度 / 安全规则自动审查"],
        ["研发", "新船型导入 Agent", "全生命周期里程碑、风险识别、批量试产"],
        ["标准", "船级社标准 Agent", "IACS / CCS / LR / DNV / ABS 规范辅助查询与判定"],
        ["研发", "设计变更 Agent", "受影响 BOM / 在制 / 库存 / 工序识别 + 协同"],
    ],
    col_widths=[2.2, 5.0, 9.0],
)

# ============ 五、生产排产决策深化（新增核心章节） ============
add_title("五、生产排产决策深化：突破造船核心瓶颈", size=15)
add_para("通用 APS / MES 在船厂「水土不服」，根因不在算力，而在它照搬离散制造逻辑，不懂造船特有的生产约束。智衍的做法不是再做一个排产引擎，而是在既有 ERP / MES 之上叠加一层「懂造船约束的实时决策脑」——把排产从「一次排定」升级为「动态再平衡」，把决策权留在人手里、把判断与预案交给分身。")

add_para("5.1 造船排产五大结构性约束", size=12.5, bold=True, color=DEEP, space_after=4)
make_table(
    ["约束维度", "造船现场特征", "通用 APS / MES 短板"],
    [
        ["船坞是唯一瓶颈", "总段合拢→下水→码头周期，坞期延误沿关键路径连锁放大", "把船坞当普通资源，缺关键路径闸门建模"],
        ["壳舾涂一体化", "先行预舾装使舾装前移，串行变并行", "照搬串行工序，不懂先行 / 后行语义"],
        ["托盘齐套耦合", "后行窝工约八成源于托盘不齐，而非工人不足", "只排工时，不排物料–工序耦合"],
        ["多项目并行", "同一船坞 / 同一批焊工同时服务 3–5 艘不同船型", "单项目优化，缺跨船型资源冲突消解"],
        ["设计变更高频", "ECN 频繁，计划脆弱", "一次排定，无动态再平衡机制"],
    ],
    col_widths=[3.4, 6.0, 6.0],
)

add_para("5.2 五类排产决策智能体", size=12.5, bold=True, color=DEEP, space_after=4)
add_para("与上述约束一一对应，智衍在造船行业 manifest 内嵌五类排产决策智能体，由既有九类智能体提供事实底座（设备健康、一次合格率、设计变更等），不重复造轮子：", space_after=4)
make_table(
    ["排产智能体", "输入", "输出 / 决策支撑", "价值落点"],
    [
        ["多项目负荷均衡 Agent", "MES 工单进度 + 设备 OEE + 工种排班", "跨船型资源冲突清单与再平衡建议", "焊工 / 船坞 / 场地利用率提升"],
        ["船坞就绪度预测 Agent", "分段完成度 + 托盘齐套率 + 涂装状态", "下水时点预测与偏差预警", "坞期延误前置拦截"],
        ["托盘齐套预警 Agent", "物料到货 + BOM + 工序计划", "后行窝工前置预警与齐套倒计时", "后行作业平滑、减少等待"],
        ["节拍偏离监控 Agent", "节拍化目标 + 实际节拍数据", "节拍偏离度与根因归因", "对标节拍化流水生产"],
        ["What-if 排产推演 Agent", "多套排产假设 + 约束参数", "改优先级→总坞期连锁影响推演", "方案比选、人留终审"],
    ],
    col_widths=[3.6, 5.0, 5.0, 2.8],
)

add_para("5.3 决策看板三件套", size=12.5, bold=True, color=DEEP, space_after=4)
make_table(
    ["看板组件", "呈现内容", "服务对象"],
    [
        ["船坞甘特", "多船型船坞占用窗口 + 关键路径高亮", "经营管控者"],
        ["负荷热力", "各车间 / 工种 / 设备实时负荷与峰谷", "生产调度主管"],
        ["就绪度仪表", "单船下水分段完成 / 托盘 / 涂装就绪度", "项目与采购负责人"],
    ],
    col_widths=[3.2, 9.2, 3.0],
)

add_para("5.4 匿名标杆测算框架（示意）", size=12.5, bold=True, color=DEEP, space_after=4)
make_table(
    ["指标", "接入前（行业经验基准）", "目标（以试点实测校验）"],
    [
        ["船坞窗口利用率", "依赖人工排程，窗口易冲突", "↑ 可量化提升（待试点实测）"],
        ["后行窝工等待", "托盘不齐致频繁等待", "↓ 齐套预警前置拦截"],
        ["计划变更响应", "按周人工重排", "实时动态再平衡"],
    ],
    col_widths=[4.0, 6.0, 5.4],
)
add_para("注：上述数值为指标框架示意，正式口径以试点实测数据与双方约定为准，本方案不作任何业绩承诺。", size=10, color=SUB)

# ============ 六、三类核心能力抓手 ============
add_title("六、三类核心能力抓手", size=15)
add_bullet("① 设备健康与预测维护：覆盖造船关键设备（船坞 / 龙门吊 / 焊接机器人 / 数控板材 / 喷砂涂装线 / 船用柴油机），事件流经 UNS 实时上屏，异常自动预警。")
add_bullet("② 质量与工艺：一次合格率 Agent 串联船体报验合格率趋势、缺陷分类与设备-质量关联根因；涂装 / 焊接缺陷判定 Agent 降低误报与复判工时；可建造性检查 Agent 把合规审查前置。")
add_bullet("③ 研发与船级社标准：新船型导入 Agent 管理全生命周期风险；船级社标准 Agent 把「翻规范」升级为「对话即得答案」；设计变更 Agent 量化变更影响并驱动跨部门协同。")
add_para("「一次合格率」是造船质量的生命线指标——智衍的目标是从被动报验走向主动预防。", size=10.5, color=SUB)

# ============ 七、技术架构 ============
add_title("七、技术架构：数字孪生五要素 + 行业 manifest", size=15)
add_para("智衍以「数字孪生五要素」为底座，构成可验证的决策闭环：")
make_table(
    ["要素", "作用"],
    [
        ["OGG 几何内核", "统一几何与拓扑真相源，支撑设备与船体结构孪生"],
        ["Agent 智能层", "九类行业智能体驱动的实时判断与预案"],
        ["实时数据", "UNS 事件流持续接入，全要素实时在线"],
        ["CAE 仿真", "工艺与结构虚拟验证，降低试错成本"],
        ["验证闭环", "结论回流资产层，持续自我进化"],
    ],
    col_widths=[4.0, 12.2],
)
add_para("行业 = manifest（设备表 / 触发词 / 智能体覆盖）。新增行业只丢一个 JSON，零代码改动——这是智衍「数据驱动、不堆代码」的核心设计。", size=10.5, color=SUB)

# ============ 八、行业验证与标杆实践（匿名） ============
add_title("八、行业验证与标杆实践（匿名）", size=15)
add_para("我们已与 3+ 家头部造船集团建立接触，覆盖三类典型形态，并沿「培训赋能 → 业务接触 → 试点进阶」路径系统化沉淀行业认知：")
make_table(
    ["客户形态（匿名）", "合作阶段", "聚焦场景"],
    [
        ["某大型造船集团（演示租户）", "试点中", "7 台关键设备点亮 + UNS 实时事件流上线，决策实时化率作为北极星被持续度量"],
        ["某邮轮与散货船建造龙头", "接触中", "专项培训与业务场景接触，聚焦设备健康与质量一次合格率"],
        ["某液化气船与高端船舶建造龙头", "培训中", "管理层培训与智能船厂蓝图探讨，聚焦研发与船级社标准"],
    ],
    col_widths=[5.0, 2.6, 8.6],
)
add_para("演示租户成效：从「看报表」升级为「实时看板 + 预警 + 预案」，经营管控者第一次拥有可对话的决策分身。", size=10.5, color=SUB)

# ============ 九、共性痛点与差异化试点 ============
add_title("九、共性痛点与差异化试点场景", size=15)
add_para("三家接触船厂虽形态各异，但共性痛点高度一致；同时基于各自业务重心，差异化试点切入点清晰：")
make_table(
    ["维度", "共性痛点", "差异化试点场景"],
    [
        ["设备", "大型专用设备停机损失高、运维靠经验", "演示租户已点亮 7 台；邮轮/散货船龙头优先设备健康"],
        ["质量", "一次合格率低、缺陷归因链条长", "全员聚焦一次合格率；涂装/焊接缺陷判定优先"],
        ["研发", "新船型与规范查询效率低", "液化气船/高端船舶龙头优先研发与船级社标准"],
        ["协同", "跨专业数据割裂、变更影响难量化", "设计变更 Agent 统一量化变更影响"],
        ["排产", "船坞瓶颈 + 托盘不齐 + 多船并行冲突", "五类排产 Agent 统一消解（见第五章）"],
    ],
    col_widths=[2.4, 6.0, 7.8],
)

# ============ 十、合作进阶路径 ============
add_title("十、合作进阶路径：G 模式三圈解锁", size=15)
add_para("智衍采用 G 模式三圈解锁，由浅入深、零风险起步，逐步迈向智能体深度合作：")
make_table(
    ["圈层", "合作模式", "交付内容", "集成要求"],
    [
        ["外圈", "免费体验", "注册即入职，零集成获得行业洞察与决策分身预生成（TTFV < 60s）", "零集成"],
        ["中圈", "接 1 源", "接入 1 路数据源（设备 / 质量 / 工艺任一），实时看板 + 预警 + 预案上线", "接 1 路数据"],
        ["内圈", "私有化", "数据本地化部署 + 全量智能体 + 行业知识资产沉淀，决策实时化率稳态达标", "私有化部署"],
    ],
    col_widths=[1.6, 2.8, 8.8, 3.0],
)

# ============ 十一、北极星指标 ============
add_title("十一、北极星指标：决策实时化率", size=15)
add_para("我们用「决策实时化率」衡量价值落地，而非传统签单口径——它直接反映经营管控者从「看报表」走向「实时决策」的比例：")
make_table(
    ["阶段", "决策实时化率", "含义"],
    [
        ["MVP 试点期", "≥ 40%", "核心场景实时看板 + 预警 + 预案可用"],
        ["稳态运营期", "≥ 85%", "全量经营决策由分身实时支撑"],
    ],
    col_widths=[3.6, 4.0, 8.6],
)

# ============ 十二、试点建议与落地下一步 ============
add_title("十二、试点建议与落地下一步", size=15)
add_para("建议分三步推进，降低决策风险、快速验证价值：")
add_bullet("第一步（0–4 周）：以「设备健康 + 质量一次合格率」双场景切入试点，接入 1 路数据源，跑通实时看板与预警。")
add_bullet("第二步（1–3 月）：扩展至研发与船级社标准，沉淀行业知识资产，决策实时化率冲刺 40%。")
add_bullet("第三步（3–6 月）：评估私有化部署，全量智能体上线，决策实时化率向 85% 稳态演进。")
add_para("排产深化可作为第二步的「高价值增量」：在设备与质量数据打通后，顺势接入 MES 工单与物料到货，即可点亮五类排产 Agent 与决策看板（详见第五章）。", size=10.5, color=SUB)
add_para("开户方式（双模）：① 名片即开户——口述公司名 / 网址即启动；② 表单注册——公司名 + 网站 + 选行业；③ 事件二维码——现场扫码即开启个性化战略。注册即入职，首屏预生成企业画像与洞察（TTFV < 60s）。", size=10.5, color=SUB)

# ============ 十三、零真名铁律与联系 ============
add_title("十三、零真名铁律声明", size=15)
add_para("本方案书对外版本严格执行零真名铁律：客户以「某大型造船集团 / 某邮轮与散货船建造龙头 / 某液化气船与高端船舶建造龙头」等形态匿名呈现，不点名、不暴露商业关系；真实主体仅限内部视图，未经授权不向任何第三方披露。", size=10.5, color=SUB)
doc.add_paragraph()
add_para("联系人：杜玉河（工业 5 点 0 产业生态联盟）", size=11, color=INK, bold=True)
add_para("体验入口：demo / EvolvIQ2026（公开演示账号）", size=11, color=INK)

doc.save(OUT)
print("SAVED:", OUT)
